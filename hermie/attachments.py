"""Attachments: file or directory paths dropped into the input box.

macOS terminals insert a dragged file as its (shell-escaped) path text, so there is no drop event to handle: the
input text is scanned for tokens that resolve to existing paths. Their content is read here, in the main process
(the sandbox only opens the workspace, and a dropped file may live anywhere), and becomes the task's *material*:
it is appended to the task text locally, goes through the same privacy gate as the task, and never leaves the
machine unredacted. Directories contribute a file tree only; the executor reads individual files when it needs them.
PDF, Word (.docx) and Excel (.xlsx) files are converted to text here (pypdf / python-docx / openpyxl) so that the
same caps and the same privacy path apply to them; other binaries are skipped.
"""
from __future__ import annotations

import fnmatch
import re
from dataclasses import dataclass, field
from pathlib import Path

# A quoted token, or a bare token in which spaces are backslash-escaped; must start like a path (/ or ~).
_PATH_RE = re.compile(r"""(?:"([^"\n]+)"|'([^'\n]+)'|((?:~|/)(?:\\ |[^\s"'])*))""")
_SKIP_DIRS = {".git", "node_modules", "__pycache__", ".venv", "venv", ".idea", ".vscode", ".DS_Store"}
_TREE_MAX_ENTRIES = 400
_NEVER_ATTACH = {Path("/"), Path.home()}   # "/" is also the slash-command prefix; neither is a sensible attachment


@dataclass
class Material:
    text: str = ""                      # what is appended to the task text (local only)
    summary: list[str] = field(default_factory=list)   # one line per attachment, for the UI
    notes: list[str] = field(default_factory=list)     # skipped / truncated items, for the UI


def find_paths(text: str) -> list[Path]:
    """Existing files/directories mentioned in the text, in order, without duplicates."""
    found: list[Path] = []
    for m in _PATH_RE.finditer(text):
        raw = next(g for g in m.groups() if g is not None)
        raw = raw.replace("\\ ", " ").rstrip(".,;:")
        if not raw.startswith(("/", "~")):
            continue
        try:
            p = Path(raw).expanduser()
            if p in _NEVER_ATTACH or not p.exists() or p in found:
                continue
            found.append(p)
        except (OSError, ValueError):
            continue
    return found


def human_size(n: int) -> str:
    for unit in ("B", "KB", "MB"):
        if n < 1024 or unit == "MB":
            return f"{n} {unit}" if unit == "B" else f"{n / 1024:.1f} {unit}"
        n //= 1024
    return f"{n} MB"


def _read_text(p: Path) -> str | None:
    """UTF-8 text, or None when the file looks binary."""
    try:
        data = p.read_bytes()
    except OSError:
        return None
    if b"\x00" in data[:4096]:
        return None
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return None


class _DocError(Exception):
    """The document could not be attached; str(e) is the note shown to the user."""


def _extract_pdf(p: Path, cap: int) -> tuple[str, str]:
    from pypdf import PdfReader
    reader = PdfReader(str(p))
    if reader.is_encrypted and not reader.decrypt(""):
        raise _DocError("encrypted, not attached")
    parts: list[str] = []
    used = 0
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(f"[Page {i}]\n{text}")
            used += len(text)
        if used >= cap:
            break
    if not parts:
        raise _DocError("no extractable text (scanned image?), not attached")
    n = len(reader.pages)
    return "\n\n".join(parts), f"{n} page{'s' if n != 1 else ''}"


def _extract_docx(p: Path, cap: int) -> tuple[str, str]:
    import docx
    from docx.table import Table
    doc = docx.Document(str(p))
    parts: list[str] = []
    used = 0
    for item in doc.iter_inner_content():  # paragraphs and tables in document order
        if isinstance(item, Table):
            rows = ["\t".join(c.text.strip() for c in row.cells) for row in item.rows]
            text = "\n".join(r for r in rows if r.strip())
        else:
            text = item.text.strip()
        if text:
            parts.append(text)
            used += len(text)
        if used >= cap:
            break
    return "\n".join(parts), "docx"


def _extract_xlsx(p: Path, cap: int) -> tuple[str, str]:
    import openpyxl
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    parts: list[str] = []
    used = 0
    n_sheets = 0
    try:
        for ws in wb.worksheets:
            rows: list[str] = []
            for row in ws.iter_rows(values_only=True):
                if all(v is None for v in row):
                    continue
                line = "\t".join("" if v is None else str(v) for v in row).rstrip("\t")
                rows.append(line)
                used += len(line)
                if used >= cap:
                    break
            if rows:
                n_sheets += 1
                parts.append(f"[Sheet: {ws.title}]\n" + "\n".join(rows))
            if used >= cap:
                break
    finally:
        wb.close()
    if not parts:
        raise _DocError("no cell data, not attached")
    return "\n\n".join(parts), f"xlsx, {n_sheets} sheet{'s' if n_sheets != 1 else ''}"


_DOC_EXTRACTORS = {".pdf": _extract_pdf, ".docx": _extract_docx, ".xlsx": _extract_xlsx, ".xlsm": _extract_xlsx}


def _extract_document(p: Path, cap: int) -> tuple[str, str]:
    """Text of a PDF / Word / Excel file (read here, outside the sandbox) and a short tag for the UI summary.
    Raises _DocError with the note to show when the file cannot be attached."""
    extract = _DOC_EXTRACTORS[p.suffix.lower()]
    try:
        return extract(p, cap)
    except _DocError:
        raise
    except Exception as e:  # corrupt file, unsupported variant, library failure
        raise _DocError(f"unreadable {p.suffix.lstrip('.')} ({type(e).__name__}), not attached") from e


def _tree(root: Path) -> tuple[list[str], int]:
    """Relative paths of the files under root (skipping VCS/dependency dirs and hidden entries), capped."""
    lines: list[str] = []
    n_files = 0
    truncated = False

    def walk(d: Path, prefix: str) -> None:
        nonlocal n_files, truncated
        try:
            entries = sorted(d.iterdir(), key=lambda e: (not e.is_dir(), e.name.lower()))
        except OSError:
            return
        for e in entries:
            if e.name in _SKIP_DIRS or e.name.startswith("."):
                continue
            if len(lines) >= _TREE_MAX_ENTRIES:
                truncated = True
                return
            rel = f"{prefix}{e.name}"
            if e.is_dir():
                lines.append(rel + "/")
                walk(e, rel + "/")
            else:
                lines.append(rel)
                n_files += 1

    walk(root, "")
    if truncated:
        lines.append(f"... (listing truncated at {_TREE_MAX_ENTRIES} entries)")
    return lines, n_files


def load_material(paths: list[Path], *, max_file_chars: int, max_total_chars: int,
                  deny_names: tuple[str, ...] = ()) -> Material:
    """Build the material block for the given attachments. Text files are included (truncated per file and in
    total); binary files, credential files and anything past the total limit are skipped with a note."""
    m = Material()
    blocks: list[str] = []
    used = 0
    for p in paths:
        if p.is_dir():
            lines, n = _tree(p)
            block = f"[Directory: {p}]\n" + "\n".join(lines)
            blocks.append(block)
            used += len(block)
            m.summary.append(f"{p}/ ({n} files)")
            continue
        if any(fnmatch.fnmatch(p.name, pat) for pat in deny_names):
            m.notes.append(f"{p.name}: credential file, not attached")
            continue
        if used >= max_total_chars:
            m.notes.append(f"{p.name}: skipped, total attachment limit reached")
            continue
        cap = min(max_file_chars, max_total_chars - used)
        if p.suffix.lower() in _DOC_EXTRACTORS:
            try:
                content, size = _extract_document(p, cap)
            except _DocError as e:
                m.notes.append(f"{p.name}: {e}")
                continue
        else:
            content = _read_text(p)
            if content is None:
                m.notes.append(f"{p.name}: binary or unreadable, not attached")
                continue
            size = human_size(len(content.encode("utf-8")))
        if len(content) > cap:
            content = content[:cap]
            m.notes.append(f"{p.name}: truncated to {cap} characters")
        block = f"[File: {p}]\n{content}"
        blocks.append(block)
        used += len(block)
        m.summary.append(f"{p} ({size})")
    m.text = "\n\n".join(blocks)
    return m
