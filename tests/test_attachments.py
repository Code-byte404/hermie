"""Attachments: paths dropped into the input box become material for the task (read in the main process, so files
outside the workspace work too; the content goes through the privacy gate like any other task text)."""
from pathlib import Path

from hermie.attachments import find_paths, load_material


def test_find_paths_recognises_existing_absolute_paths(tmp_path):
    f = tmp_path / "notes.txt"
    f.write_text("hi")
    text = f"summarize {f} please"
    assert find_paths(text) == [f]


def test_find_paths_handles_escaped_spaces_quotes_and_tilde(tmp_path, monkeypatch):
    d = tmp_path / "my docs"
    d.mkdir()
    f = d / "a b.md"
    f.write_text("x")
    monkeypatch.setenv("HOME", str(tmp_path))
    escaped = str(f).replace(" ", "\\ ")
    quoted = f"'{f}'"
    tilde = "~/my\\ docs/a\\ b.md"
    assert find_paths(f"read {escaped}") == [f]
    assert find_paths(f"read {quoted}") == [f]
    assert find_paths(f"read {tilde}") == [f]


def test_find_paths_unescapes_shell_metacharacters(tmp_path):
    # What Terminal / iTerm2 / Ghostty insert for a dropped "Q3 report (final) & notes's.pdf".
    f = tmp_path / "Q3 report (final) & notes's.pdf"
    f.write_text("x")
    escaped = str(f).replace(" ", "\\ ").replace("(", "\\(").replace(")", "\\)").replace("&", "\\&").replace("'", "\\'")
    assert find_paths(f"summarize {escaped} ") == [f]
    assert find_paths(f"summarize{escaped}") == [f]   # dropped right after a word, no space


def test_find_paths_ignores_missing_paths_and_dedupes(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("x")
    text = f"{f} and /nonexistent/zzz.txt and {f} again"
    assert find_paths(text) == [f]


def test_load_material_reads_text_files_with_headers(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("hello world")
    m = load_material([f], max_file_chars=1000, max_total_chars=5000)
    assert f"[File: {f}]" in m.text and "hello world" in m.text
    assert m.summary == [f"{f} (11 B)"]
    assert m.notes == []


def test_load_material_skips_binary_and_truncates_large(tmp_path):
    b = tmp_path / "img.png"
    b.write_bytes(b"\x89PNG\x00\x00binary")
    big = tmp_path / "big.txt"
    big.write_text("x" * 500)
    m = load_material([b, big], max_file_chars=100, max_total_chars=5000)
    assert "binary" not in m.text
    assert any("img.png" in n and "binary" in n for n in m.notes)
    assert m.text.split("]\n", 1)[1] == "x" * 100
    assert any("big.txt" in n and "truncated" in n for n in m.notes)


def test_load_material_directory_gives_tree_only(tmp_path):
    d = tmp_path / "proj"
    (d / "src").mkdir(parents=True)
    (d / "src" / "main.py").write_text("SECRET_CONTENT = 1")
    (d / "README.md").write_text("readme body")
    (d / ".git").mkdir()
    (d / ".git" / "HEAD").write_text("ref")
    (d / "node_modules").mkdir()
    (d / "node_modules" / "x.js").write_text("js")
    m = load_material([d], max_file_chars=1000, max_total_chars=5000)
    assert f"[Directory: {d}]" in m.text
    assert "src/main.py" in m.text and "README.md" in m.text
    assert "SECRET_CONTENT" not in m.text and "readme body" not in m.text
    assert ".git" not in m.text and "node_modules" not in m.text
    assert m.summary == [f"{d}/ (2 files)"]


def test_load_material_directory_says_whether_tools_can_open_it(tmp_path):
    ws = tmp_path / "ws"
    inside = ws / "pkg"
    outside = tmp_path / "elsewhere"
    for d in (inside, outside):
        d.mkdir(parents=True)
        (d / "a.py").write_text("x")
    m = load_material([inside, outside], max_file_chars=1000, max_total_chars=5000, workspace=ws)
    assert "Inside the workspace at pkg/" in m.text
    assert "Outside the workspace, attached read-only" in m.text
    assert m.roots == [inside, outside]


def test_load_material_refuses_credential_files(tmp_path):
    k = tmp_path / "id_rsa"
    k.write_text("-----BEGIN OPENSSH PRIVATE KEY-----\nabc")
    m = load_material([k], max_file_chars=1000, max_total_chars=5000, deny_names=("id_rsa", "*.pem"))
    assert "PRIVATE KEY" not in m.text
    assert any("id_rsa" in n and "credential" in n for n in m.notes)


def test_load_material_total_cap_stops_further_files(tmp_path):
    a, b = tmp_path / "a.txt", tmp_path / "b.txt"
    a.write_text("a" * 80)
    b.write_text("b" * 80)
    m = load_material([a, b], max_file_chars=1000, max_total_chars=100)
    assert "a" * 80 in m.text
    assert "b" * 80 not in m.text
    assert any("b.txt" in n and "limit" in n for n in m.notes)


def test_find_paths_ignores_root_and_home():
    assert find_paths("/") == []
    assert find_paths("/help") == []
    assert find_paths("~") == []
    assert find_paths(f"list {Path.home()}") == []


# ---- office / PDF documents: text is extracted in the main process and joins the material like any text file ----

def _minimal_pdf(pages: list[str]) -> bytes:
    """A hand-built single-font PDF with one line of text per page (pypdf can read but not author text)."""
    objs: list[bytes] = []
    n_pages = len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n_pages))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    font_id = 3 + 2 * n_pages
    for i, line in enumerate(pages):
        page_id, content_id = 3 + 2 * i, 4 + 2 * i
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents {content_id} 0 R "
                    f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode())
        stream = f"BT /F1 12 Tf 20 100 Td ({line}) Tj ET".encode()
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def test_load_material_extracts_pdf_text_per_page(tmp_path):
    f = tmp_path / "report.pdf"
    f.write_bytes(_minimal_pdf(["Revenue grew 12 percent", "Phone 13812345678"]))
    m = load_material([f], max_file_chars=10_000, max_total_chars=50_000)
    assert f"[File: {f}]" in m.text
    assert "Revenue grew 12 percent" in m.text and "13812345678" in m.text
    assert m.text.index("[Page 1]") < m.text.index("Revenue") < m.text.index("[Page 2]") < m.text.index("13812345678")
    assert m.summary == [f"{f} (2 pages)"]
    assert m.notes == []


def test_load_material_notes_pdf_without_text_and_encrypted(tmp_path):
    from pypdf import PdfReader, PdfWriter
    blank = tmp_path / "scan.pdf"
    w = PdfWriter()
    w.add_blank_page(width=100, height=100)
    with blank.open("wb") as fh:
        w.write(fh)
    locked = tmp_path / "locked.pdf"
    w2 = PdfWriter(clone_from=PdfReader(str(blank)))
    w2.encrypt("secret")
    with locked.open("wb") as fh:
        w2.write(fh)
    m = load_material([blank, locked], max_file_chars=10_000, max_total_chars=50_000)
    assert m.text == "" and m.summary == []
    assert any("scan.pdf" in n and "no extractable text" in n for n in m.notes)
    assert any("locked.pdf" in n and "encrypted" in n for n in m.notes)


def test_load_material_extracts_docx_paragraphs_and_tables(tmp_path):
    import docx
    d = docx.Document()
    d.add_heading("Quarterly review", level=1)
    d.add_paragraph("Customer Zhang called about invoice 42.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "name", "phone"
    t.cell(1, 0).text, t.cell(1, 1).text = "Li", "13912345678"
    f = tmp_path / "review.docx"
    d.save(str(f))
    m = load_material([f], max_file_chars=10_000, max_total_chars=50_000)
    assert "Quarterly review" in m.text and "invoice 42" in m.text
    assert "Li\t13912345678" in m.text
    assert m.summary == [f"{f} (docx)"]
    assert m.notes == []


def test_load_material_extracts_xlsx_sheets_as_tab_separated_rows(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Customers"
    ws.append(["name", "phone", "amount"])
    ws.append(["Wang", "13712345678", 99.5])
    ws2 = wb.create_sheet("Empty")
    f = tmp_path / "book.xlsx"
    wb.save(str(f))
    m = load_material([f], max_file_chars=10_000, max_total_chars=50_000)
    assert "[Sheet: Customers]" in m.text
    assert "name\tphone\tamount" in m.text and "Wang\t13712345678\t99.5" in m.text
    assert "[Sheet: Empty]" not in m.text
    assert m.summary == [f"{f} (xlsx, 1 sheet)"]


def test_load_material_document_text_respects_caps(tmp_path):
    import docx
    d = docx.Document()
    for _ in range(50):
        d.add_paragraph("word " * 20)
    f = tmp_path / "long.docx"
    d.save(str(f))
    m = load_material([f], max_file_chars=200, max_total_chars=50_000)
    assert len(m.text.split("]\n", 1)[1]) == 200
    assert any("long.docx" in n and "truncated" in n for n in m.notes)


def test_load_material_corrupt_document_is_skipped_with_note(tmp_path):
    f = tmp_path / "broken.docx"
    f.write_bytes(b"PK\x03\x04 not really a docx")
    m = load_material([f], max_file_chars=1000, max_total_chars=5000)
    assert m.text == ""
    assert any("broken.docx" in n and "unreadable" in n for n in m.notes)
